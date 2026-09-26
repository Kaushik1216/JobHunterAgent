from __future__ import annotations

import json
import re
import uuid
from pathlib import Path
from typing import Any

import structlog

from inference.client import LLMClient
from inference.prompts import (
    RESUME_EXTRACTION_SYSTEM_PROMPT,
    build_resume_extraction_prompt,
)
from models.schemas import (
    CategorizedSkills,
    EducationItem,
    ParsedResumeProfile,
    ProjectItem,
    ResumeRecord,
    WorkExperienceItem,
)

logger = structlog.get_logger(__name__)


def extract_text_from_file(file_path: Path) -> str:
    """Extract raw text from PDF, DOCX, TXT, or Markdown resume files."""
    suffix = file_path.suffix.lower()
    
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(file_path))
            pages_text = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text)
            text = "\n\n".join(pages_text)
            if text.strip():
                return _clean_extracted_text(text)
        except Exception as e:
            logger.warning("pypdf_extraction_failed", error=str(e), path=str(file_path))

    elif suffix in (".docx", ".doc"):
        try:
            import docx
            doc = docx.Document(str(file_path))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        lines.append(row_text)
            text = "\n".join(lines)
            if text.strip():
                return _clean_extracted_text(text)
        except Exception as e:
            logger.warning("docx_extraction_failed", error=str(e), path=str(file_path))

    # Fallback to UTF-8 text read (works for .txt, .md, and plain text files)
    try:
        return _clean_extracted_text(file_path.read_text(encoding="utf-8", errors="ignore"))
    except Exception as e:
        logger.error("text_file_read_failed", error=str(e), path=str(file_path))
        return ""


def _clean_extracted_text(text: str) -> str:
    """Normalize whitespace and remove excessive blank lines."""
    # Replace multiple spaces with a single space
    cleaned = re.sub(r"[ \t]+", " ", text)
    # Replace 3 or more newlines with double newlines
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def extract_json_from_llm(text: str) -> dict[str, Any]:
    """Extract valid JSON from raw LLM responses including code fences or raw text."""
    if not text:
        return {}

    # 1. Match code fence ```json ... ``` or ``` ... ```
    fence_pattern = r"```(?:json)?\s*([\s\S]*?)\s*```"
    match = re.search(fence_pattern, text)
    if match:
        content = match.group(1).strip()
        try:
            return json.loads(content)
        except json.JSONDecodeError:
            pass

    # 2. Match outermost { ... }
    brace_pattern = r"(\{[\s\S]*\})"
    match = re.search(brace_pattern, text)
    if match:
        content = match.group(1).strip()
        try:
            return json.loads(content)
        except json.JSONDecodeError:
            pass

    # 3. Direct JSON load attempt
    try:
        return json.loads(text.strip())
    except json.JSONDecodeError:
        logger.error("failed_to_extract_json_from_llm", raw=text[:300])
        return {}


def parse_resume_with_ai(raw_text: str, client: LLMClient) -> ParsedResumeProfile:
    """Use configured SLM/LLM to extract structured profile information from resume."""
    if not raw_text or len(raw_text.strip()) < 50:
        logger.warning("resume_text_too_short_for_ai", length=len(raw_text))
        return ParsedResumeProfile(summary="Resume text too short or empty.")

    prompt = build_resume_extraction_prompt(raw_text)
    system_prompt = RESUME_EXTRACTION_SYSTEM_PROMPT

    logger.info("requesting_ai_resume_extraction", model=client.model, provider=client.provider)
    
    try:
        response_text = client.generate(prompt=prompt, system_prompt=system_prompt)
        parsed_dict = extract_json_from_llm(response_text)
        
        if not parsed_dict:
            # Retry once with explicit json reminder
            retry_prompt = f"Please return ONLY the valid JSON object without any explanations:\n\n{response_text}"
            retry_text = client.generate(prompt=retry_prompt, system_prompt="You are a JSON fixer. Return only valid JSON.")
            parsed_dict = extract_json_from_llm(retry_text)

        if not parsed_dict:
            logger.error("resume_ai_parsing_returned_empty_dict")
            return _create_fallback_profile(raw_text)

        # Ensure all_skills contains flat list
        skills_data = parsed_dict.get("skills", {})
        if isinstance(skills_data, dict):
            all_collected = []
            for key in ("languages", "frameworks_and_libraries", "databases", "cloud_and_devops", "tools_and_platforms", "core_competencies"):
                items = skills_data.get(key, [])
                if isinstance(items, list):
                    all_collected.extend(items)
            existing_all = skills_data.get("all_skills", [])
            if isinstance(existing_all, list) and existing_all:
                all_collected.extend(existing_all)
            # Deduplicate while preserving order
            seen = set()
            skills_data["all_skills"] = [s for s in all_collected if isinstance(s, str) and not (s.lower() in seen or seen.add(s.lower()))]
            parsed_dict["skills"] = skills_data

        return ParsedResumeProfile.model_validate(parsed_dict)
    except Exception as e:
        logger.error("resume_ai_extraction_exception", error=str(e))
        return _create_fallback_profile(raw_text)


def _create_fallback_profile(raw_text: str) -> ParsedResumeProfile:
    """Basic rule-based fallback if AI inference is unavailable or fails."""
    # Simple regex extraction for email and phone
    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", raw_text)
    phone_match = re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", raw_text)

    # Basic skill keywords heuristic
    common_skills = [
        "Python", "JavaScript", "TypeScript", "React", "Node.js", "FastAPI", "SQL",
        "PostgreSQL", "Docker", "AWS", "Git", "Linux", "Kubernetes", "Java", "C++",
        "MongoDB", "Redis", "HTML", "CSS", "REST", "CI/CD",
    ]
    detected_skills = [s for s in common_skills if re.search(rf"\b{re.escape(s)}\b", raw_text, re.IGNORECASE)]

    return ParsedResumeProfile(
        name=None,
        email=email_match.group(0) if email_match else None,
        phone=phone_match.group(0) if phone_match else None,
        summary=raw_text[:500] + ("..." if len(raw_text) > 500 else ""),
        skills=CategorizedSkills(all_skills=detected_skills),
        target_roles=["Software Engineer"],
    )


def process_resume_upload(
    file_bytes: bytes,
    original_filename: str,
    resumes_dir: Path,
    client: LLMClient | None = None,
) -> ResumeRecord:
    """Save resume file to project directory, extract text, and extract structured profile via AI."""
    resumes_dir.mkdir(parents=True, exist_ok=True)
    resume_id = str(uuid.uuid4())
    safe_name = re.sub(r"[^\w\-.]", "_", original_filename)
    dest_path = resumes_dir / f"{resume_id}_{safe_name}"
    
    # Save file to disk
    dest_path.write_bytes(file_bytes)
    logger.info("resume_file_persisted", path=str(dest_path), size=len(file_bytes))

    # Extract text
    raw_text = extract_text_from_file(dest_path)
    file_type = dest_path.suffix.lower().lstrip(".")

    # AI structured extraction (One-time job)
    parsed_profile = None
    if client is not None and raw_text:
        parsed_profile = parse_resume_with_ai(raw_text, client)

    candidate_name = parsed_profile.name if parsed_profile else None
    headline = parsed_profile.headline if parsed_profile else None
    email = parsed_profile.email if parsed_profile else None
    total_yoe = parsed_profile.total_experience_years if parsed_profile else None

    return ResumeRecord(
        id=resume_id,
        filename=original_filename,
        file_path=str(dest_path),
        file_type=file_type,
        raw_text=raw_text,
        parsed_profile=parsed_profile,
        candidate_name=candidate_name,
        headline=headline,
        email=email,
        total_yoe=total_yoe,
        is_active=True,
    )
